"""Build the Word companion from the reviewed Markdown; use python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
source = (root/'docs/LEARNING-MANUAL.md').read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.7)
section.page_width, section.page_height = Inches(8.5), Inches(11)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for fonts in style.element.iter(qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if attr.endswith('Theme'):del fonts.attrib[attr]
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(28)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12.5)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(14)
    doc.styles[name].paragraph_format.space_after=Pt(7)
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Ethical Walls Matter  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def inline(paragraph, text):
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
        run=paragraph.add_run(part.strip('`') if part.startswith('`') else part.strip('*') if part.startswith('**') else part)
        if part.startswith('`'):
            run.font.name='Consolas';run.font.size=Pt(9)
        elif part.startswith('**'):run.bold=True

def table(rows):
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=([2.1,5.0] if len(rows[0])==2 else [1.8,2.3,3.0])
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for n,row in enumerate(rows):
        cells=t.rows[0].cells if n==0 else t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            inline(cell.paragraphs[0],text)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for run in p.runs:run.font.size=Pt(9);run.font.name='Calibri'
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2' if n==0 else 'F5F7F9' if n%2==0 else 'FFFFFF');tcPr.append(shade)
            if n==0:
                for run in cell.paragraphs[0].runs:run.bold=True
        if n==0:
            header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=source.splitlines();i=0;code=False
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        code=not code;i+=1;continue
    if code:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
        p.paragraph_format.left_indent=Inches(.08)
        r=p.add_run(line or ' ');r.font.name='Consolas';r.font.size=Pt(8)
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            parts=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x) for x in parts):rows.append(parts)
            i+=1
        table(rows);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line:
        p=doc.add_paragraph();inline(p,line)
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.left_indent=Inches(.18);p.paragraph_format.first_line_indent=Inches(-.18)
    i+=1
doc.core_properties.title='Ethical Walls Matter Learning Manual'
doc.core_properties.subject='SQL views service accounts and local OAuth authentication'
doc.core_properties.author='Ethical Walls Matter'
out=root/'docs/Ethical-Walls-Matter-Learning-Manual.docx'
doc.save(out)
print(out)
